"""Importa las hojas 'Consolidado Contratos' y 'Consolidado Convenios' a PostgreSQL.

Uso:
    python -m importer.importar RUTA_EXCEL.xlsx [--reporte reporte_importacion.xlsx]

Borra y vuelve a cargar todas las tablas: se puede correr cuantas veces haga falta
hasta el corte definitivo. Genera un Excel de reporte con:
  - Problemas: celdas que no se pudieron interpretar (fila, columna, valor, motivo)
  - Personas: nombres encontrados, sus variantes de escritura y cuántos registros tienen
  - Catálogos: estados, equipos y fuentes de financiación con sus variantes
  - Resumen: conteos de lo cargado
"""
import argparse
import re
import sys
import warnings
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill

from app.db import Base, SessionLocal, engine
from app import models as m
from importer import limpieza as L

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

HOJA_CT = "Consolidado Contratos"
HOJA_CV = "Consolidado Convenios"
ORDINALES = ["PRIMER", "SEGUNDO", "TERCER", "CUARTO", "QUINTO", "SEXTO",
             "SEPTIMO", "OCTAVO", "NOVENO", "DÉCIMO"]


def norm_header(h):
    return L.clave(str(h or ""))


def verificar(encabezados, esperado: dict, hoja):
    """Comprueba que cada índice tenga el encabezado esperado (protege contra columnas movidas)."""
    errores = []
    for idx, fragmento in esperado.items():
        real = norm_header(encabezados[idx]) if idx < len(encabezados) else ""
        if L.clave(fragmento) not in real:
            errores.append(f"  col {idx + 1}: se esperaba '{fragmento}', hay '{encabezados[idx] if idx < len(encabezados) else None}'")
    if errores:
        sys.exit(f"La hoja '{hoja}' no tiene la estructura esperada:\n" + "\n".join(errores))


def filas(ws):
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue
        if any(c not in (None, "") for c in r):
            yield i, r


class Catalogos:
    """Normaliza estados, equipos y personas; guarda las variantes para el reporte."""

    def __init__(self, s):
        self.s = s
        self.estados, self.equipos, self.personas = {}, {}, {}
        self.variantes = defaultdict(Counter)  # (tipo, clave) -> Counter(texto original)

    def _get(self, cache, tipo, modelo, v, **extra):
        t = L.texto(v)
        if L.es_vacio(t):
            return None
        k = L.clave(t)
        self.variantes[(tipo, k)][t] += 1
        if k not in cache:
            obj = modelo(nombre=re.sub(r"\s+", " ", t), **extra)
            self.s.add(obj)
            self.s.flush()
            cache[k] = obj
        return cache[k]

    def estado(self, v):
        return self._get(self.estados, "estado", m.Estado, v)

    def equipo(self, v):
        return self._get(self.equipos, "equipo", m.Equipo, v)

    def persona(self, v):
        if not L.es_persona(v):
            return None
        t = L.texto(v)
        return self._get(self.personas, "persona", m.Persona, v, clave=L.clave(t))

    def nombres_canonicos(self):
        """Usa como nombre la variante más frecuente."""
        for tipo, cache in (("estado", self.estados), ("equipo", self.equipos), ("persona", self.personas)):
            for k, obj in cache.items():
                obj.nombre = re.sub(r"\s+", " ", self.variantes[(tipo, k)].most_common(1)[0][0])


def asignar(cat, obj, r, idx_sup, idx_tec, idx_fin, uso_personas):
    for idx, rol in ((idx_sup, "supervisor"), (idx_tec, "apoyo_tecnico"), (idx_fin, "apoyo_financiero")):
        p = cat.persona(r[idx])
        if p:
            obj.asignaciones.append(m.Asignacion(persona=p, rol=rol, texto_original=L.texto(r[idx])))
            uso_personas[(p.clave, rol)] += 1
        elif not L.es_vacio(r[idx]):
            uso_personas[("(sin persona) " + L.clave(str(r[idx])), rol)] += 1


def informe_si_hay_datos(**campos):
    # Un informe existe si tiene al menos una fecha, radicado, monto o enlace.
    # Las observaciones sueltas (p. ej. "Este contrato solo debe entregar 3 informes")
    # quedan en datos_originales del contrato.
    datos = {k: v for k, v in campos.items() if k not in ("tipo", "numero", "observacion")}
    if any(v not in (None, "") for v in datos.values()):
        return m.Informe(**campos)
    return None


# --------------------------------------------------------------------- contratos
CT_ESPERADO = {
    1: "CONTRATO / CONVENIO", 2: "No CT", 3: "EQUIPO", 4: "SUPERVISOR",
    5: "APOYO SUPERVISIÓN TÉCNICO", 6: "APOYO SUPERVISIÓN FINANCIERO", 9: "OBJETO",
    11: "ESTADO", 12: "ALERTAS", 35: "FUENTE FINANCIACIÓN", 50: "MONTO SOLICITADO",
    60: "MONTO DESEMBOLSO 1", 77: "FECHA ENTREGA PROGRAMADA PRIMER INFORME",
    167: "FECHA ENTREGA PROGRAMADA DÉCIMO INFORME", 177: "FECHA DE SOLICITUD DE MODIFICACIÓN",
    183: "FECHA ENTREGA INFORME FINAL PROGRAMADA", 192: "AJUSTADO 1", 196: "AJUSTADO 2",
    200: "MONTO A REINTEGRAR", 209: "PRODUCTOS VERIFICADOS",
}
# Informes de avance: 10 bloques de 10 columnas desde la 77
CT_INF = ["fecha_programada", "fecha_efectiva", "fecha_generacion_tecnico", "fecha_generacion_financiero",
          "fecha_generacion_integral", "fecha_supervision", "radicado", "monto_aprobado", "enlace", "observacion"]


def importar_contratos(ws, s, cat, log, uso_personas, encabezados):
    verificar(encabezados, CT_ESPERADO, HOJA_CT)
    for n, ordinal in enumerate(ORDINALES):
        verificar(encabezados, {77 + 10 * n: f"PROGRAMADA {ordinal} INFORME"}, HOJA_CT)
    total = 0
    for fila, r in filas(ws):
        r = list(r) + [None] * (211 - len(r))
        numero = L.texto(r[2])
        H = lambda i: str(encabezados[i]).strip()
        ctx = lambda i: (HOJA_CT, fila, numero, H(i))
        F = lambda i: L.fecha(r[i], log, ctx(i))
        M = lambda i: L.monto(r[i], log, ctx(i))
        T = lambda i: L.texto(r[i])

        if not numero:
            log.anotar(HOJA_CT, fila, None, H(2), r[2], "registro sin número de contrato")
        elif not L.NUMERO_VALIDO.match(numero):
            log.anotar(HOJA_CT, fila, numero, H(2), r[2], "número con formato distinto a XXX-AAAA")

        c = m.Contrato(
            fila_excel=fila, tipo_registro=T(1), numero=numero, fuente_financiacion=T(35),
            equipo=cat.equipo(r[3]), estado=cat.estado(r[11]),
            id_mgi=L.codigo(r[7]), codigo_sigp=L.codigo(r[8]), objeto=T(9), entidad_ejecutora=T(10),
            tipo_derivado=T(14), fecha_inicio=F(15), fecha_finalizacion=F(16), tiempo_ejecucion=T(17),
            avance_tecnico=L.porcentaje(r[18], log, ctx(18)), avance_financiero=L.porcentaje(r[19], log, ctx(19)),
            mecanismo=T(20), numero_convocatoria=T(21), titulo_convocatoria=T(22), complemento_mecanismo=T(23),
            expediente_virtual=L.codigo(r[24]), departamento=T(25), municipio=T(26), entidades_coejecutoras=T(27),
            investigador_principal=T(28), email_investigador=T(29), beneficiarios=T(30),
            participacion_minorias=T(31), descripcion_minorias=T(32), fecha_legalizacion=F(33),
            fecha_vigencia_poliza=F(34), convenio_del_que_deriva=T(36), fecha_vencimiento_convenio=F(37),
            monto_solicitado=M(50), monto_contrapartida=M(51), monto_seguimiento=M(52),
            fecha_reintegro_seguimiento=F(53), costo=M(54), monto_a_desembolsar=M(55),
            desembolsos_programados=T(56), desembolsos_realizados=T(57),
            monto_desembolsado_total=M(58), monto_por_desembolsar=M(59),
            monto_reintegrar=M(200), fecha_efectiva_reintegro=F(201), fecha_memo_liquidacion=F(202),
            radicado_memo_liquidacion=L.codigo(r[203]), enlace_memo_liquidacion=T(204),
            fecha_acta_liquidacion=F(205), fecha_notificacion_liquidacion=F(206), observaciones=T(207),
            productos_comprometidos=T(208), productos_verificados=T(209),
            datos_originales={str(encabezados[i]).strip(): L.texto(v)
                              for i, v in enumerate(r[1:210], start=1) if v not in (None, "")},
        )
        asignar(cat, c, r, 4, 5, 6, uso_personas)

        # CDR (4 bloques: convenio, CDR, valor)
        for k in range(4):
            b = 38 + 3 * k
            if any(not L.es_vacio(r[b + j]) for j in range(3)):
                c.fuentes_cdr.append(m.FuenteCDR(orden=k + 1, convenio=T(b), cdr=L.codigo(r[b + 1]), valor=M(b + 2)))

        # Desembolsos: el 1 tiene monto+fecha; del 2 al 6, monto+programada+efectiva
        if not (L.es_vacio(r[60]) and L.es_vacio(r[61])):
            c.desembolsos.append(m.Desembolso(numero=1, monto=M(60), fecha_efectiva=F(61)))
        for k in range(2, 7):
            b = 62 + 3 * (k - 2)
            if any(not L.es_vacio(r[b + j]) for j in range(3)):
                c.desembolsos.append(m.Desembolso(numero=k, monto=M(b), fecha_programada=F(b + 1), fecha_efectiva=F(b + 2)))

        # Informes de avance
        for k in range(10):
            b = 77 + 10 * k
            vals = {}
            for j, campo in enumerate(CT_INF):
                i = b + j
                if campo.startswith("fecha"):
                    vals[campo] = F(i)
                elif campo == "monto_aprobado":
                    vals[campo] = M(i)
                elif campo == "radicado":
                    vals[campo] = L.codigo(r[i])
                else:
                    vals[campo] = T(i)
            inf = informe_si_hay_datos(tipo="avance", numero=k + 1, **vals)
            if inf:
                c.informes.append(inf)

        # Informe final
        inf = informe_si_hay_datos(
            tipo="final", numero=1, fecha_programada=F(183), fecha_efectiva=F(184),
            fecha_generacion_tecnico=F(185), fecha_generacion_financiero=F(186), fecha_generacion_integral=F(187),
            fecha_supervision=F(188), radicado=L.codigo(r[189]), monto_aprobado=M(190), enlace=T(191))
        if inf:
            c.informes.append(inf)
        for k, b in ((1, 192), (2, 196)):
            inf = informe_si_hay_datos(tipo="final_ajustado", numero=k, fecha_supervision=F(b),
                                       radicado=L.codigo(r[b + 1]), monto_aprobado=M(b + 2), enlace=T(b + 3))
            if inf:
                c.informes.append(inf)

        # Modificaciones
        if any(not L.es_vacio(r[i]) for i in range(177, 183)):
            c.modificaciones.append(m.Modificacion(
                fecha_solicitud_entidad=F(177), fecha_solicitud_area_tecnica=F(178), condicion_otrosi=T(179),
                tiempo_adicion=T(180), fecha_legalizacion=F(181), numero_modificaciones=T(182)))

        # Alertas
        if not L.es_vacio(r[12]) or not L.es_vacio(r[13]):
            c.alertas.append(m.Alerta(tipo=T(12), observacion=T(13), abierta=True))

        s.add(c)
        total += 1
    return total


# --------------------------------------------------------------------- convenios
CV_ESPERADO = {
    0: "ID CV", 1: "No CV", 2: "EQUIPO", 3: "SUPERVISOR", 4: "APOYO SUPERVISIÓN TÉCNICO",
    5: "APOYO SUPERVISIÓN FINANCIERO", 6: "TIPO DE CONVENIO", 8: "ESTADO", 13: "ENTIDAD ALIADA 1",
    24: "VALOR DEL CONVENIO", 34: "PERIODICIDAD", 35: "FECHA ENTREGA PROGRAMADA INFORME 1",
    125: "FECHA ENTREGA PROGRAMADA INFORME FINAL", 134: "FECHA MEMO LIQUIDACIÓN",
    138: "COMITÉ SEGUN CONVENIO", 152: "SUPERVISOR POR MINUTA",
}
CV_INF = ["fecha_programada", "fecha_efectiva", "fecha_generacion_tecnico", "fecha_generacion_financiero",
          "fecha_generacion_integral", "fecha_supervision", "radicado", "monto_aprobado", "enlace"]


def importar_convenios(ws, s, cat, log, uso_personas, encabezados):
    verificar(encabezados, CV_ESPERADO, HOJA_CV)
    for k in range(10):
        verificar(encabezados, {35 + 9 * k: f"PROGRAMADA INFORME {k + 1}"}, HOJA_CV)
    total = 0
    for fila, r in filas(ws):
        r = list(r) + [None] * (163 - len(r))
        numero = L.texto(r[1])
        H = lambda i: str(encabezados[i]).strip()
        ctx = lambda i: (HOJA_CV, fila, numero, H(i))
        F = lambda i: L.fecha(r[i], log, ctx(i))
        M = lambda i: L.monto(r[i], log, ctx(i))
        T = lambda i: L.texto(r[i])
        if not numero:
            log.anotar(HOJA_CV, fila, None, H(1), r[1], "registro sin número de convenio")

        cv = m.Convenio(
            fila_excel=fila, id_cv=L.codigo(r[0]), numero=numero, equipo=cat.equipo(r[2]), estado=cat.estado(r[8]),
            tipo_convenio=T(6), objeto=T(7), fecha_inicio=F(10), fecha_finalizacion=F(11),
            expediente_virtual=L.codigo(r[12]), bpin_1=L.codigo(r[22]), bpin_2=L.codigo(r[23]),
            valor=M(24), monto_pagado=M(25), monto_pendiente=M(26), rendimientos=M(27),
            periodicidad_informes_meses=L.codigo(r[34]),
            fecha_memo_liquidacion=F(134), radicado_memo_liquidacion=L.codigo(r[135]),
            enlace_memo_liquidacion=T(136), fecha_acta_liquidacion=F(137),
            comite_segun_convenio=T(138), tipo_comite_1=T(139), delegado_comite_1=T(140),
            tipo_comite_2=T(141), delegado_comite_2=T(142), conformacion_comite=T(143),
            contacto_notificaciones=T(144), periodicidad_comite=T(145), fecha_ultima_reunion=F(146),
            proxima_reunion=T(147), avance_tecnico=T(148), fecha_evaluacion_avance_tecnico=T(149),
            avance_financiero=T(150), fecha_evaluacion_avance_financiero=T(151), supervisor_por_minuta=T(152),
            datos_originales={str(encabezados[i]).strip(): L.texto(v)
                              for i, v in enumerate(r[:153]) if v not in (None, "")},
        )
        asignar(cat, cv, r, 3, 4, 5, uso_personas)

        for k in range(3):
            b = 13 + 3 * k
            if any(not L.es_vacio(r[b + j]) for j in range(3)):
                cv.aliados.append(m.ConvenioAliado(orden=k + 1, entidad=T(b), aporte_efectivo=M(b + 1), aporte_especie=M(b + 2)))

        for k in range(11):  # 10 de avance + final
            b = 35 + 9 * k
            vals = {}
            for j, campo in enumerate(CV_INF):
                i = b + j
                if campo.startswith("fecha"):
                    vals[campo] = F(i)
                elif campo == "monto_aprobado":
                    vals[campo] = M(i)
                elif campo == "radicado":
                    vals[campo] = L.codigo(r[i])
                else:
                    vals[campo] = T(i)
            inf = informe_si_hay_datos(tipo="final" if k == 10 else "avance", numero=1 if k == 10 else k + 1, **vals)
            if inf:
                cv.informes.append(inf)

        if any(not L.es_vacio(r[i]) for i in range(28, 34)):
            cv.modificaciones.append(m.Modificacion(
                fecha_solicitud_entidad=F(28), fecha_solicitud_area_tecnica=F(29), condicion_otrosi=T(30),
                tiempo_adicion=T(31), fecha_legalizacion=F(32), numero_modificaciones=T(33)))

        if not L.es_vacio(r[9]):
            cv.alertas.append(m.Alerta(tipo=T(9), abierta=True))

        s.add(cv)
        total += 1
    return total


# --------------------------------------------------------------------- reporte
def escribir_reporte(ruta, log, cat, uso_personas, resumen, fuentes):
    wb = openpyxl.Workbook()
    neg = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F4E79")

    def hoja(titulo, cab, datos, anchos):
        ws = wb.create_sheet(titulo)
        ws.append(cab)
        for c in ws[1]:
            c.font, c.fill = neg, relleno
        for d in datos:
            ws.append(d)
        for i, w in enumerate(anchos):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    hoja("Resumen", ["Concepto", "Cantidad"], resumen, [45, 12])

    agrupado = Counter((p["hoja"], p["columna"], p["problema"]) for p in log.problemas)
    hoja("Problemas por columna", ["Hoja", "Columna", "Problema", "Celdas"],
         [list(k) + [v] for k, v in agrupado.most_common()], [24, 55, 45, 10])
    hoja("Problemas detalle", ["Hoja", "Fila Excel", "Número", "Columna", "Valor original", "Problema"],
         [[p["hoja"], p["fila_excel"], p["numero"], p["columna"], p["valor_original"], p["problema"]]
          for p in log.problemas], [24, 10, 16, 50, 45, 40])

    filas_p = []
    for k, p in sorted(cat.personas.items(), key=lambda x: x[1].nombre):
        variantes = cat.variantes[("persona", k)]
        filas_p.append([p.nombre, " | ".join(f"{t} ({n})" for t, n in variantes.most_common()),
                        uso_personas[(k, "supervisor")], uso_personas[(k, "apoyo_tecnico")],
                        uso_personas[(k, "apoyo_financiero")], "", "", ""])
    hoja("Personas", ["Nombre (variante más usada)", "Variantes encontradas (registros)", "Como supervisor",
                      "Como apoyo técnico", "Como apoyo financiero", "Unir con (nombre)", "Correo institucional", "Equipo"],
         filas_p, [40, 70, 14, 16, 18, 30, 34, 30])

    sin_persona = Counter()
    for (k, rol), n in uso_personas.items():
        if k.startswith("(sin persona) "):
            sin_persona[(k.replace("(sin persona) ", ""), rol)] += n
    hoja("Valores que no son persona", ["Valor", "Columna", "Registros"],
         [[k, rol, n] for (k, rol), n in sin_persona.most_common()], [50, 20, 10])

    cats = []
    for tipo in ("estado", "equipo"):
        for (t, k), var in sorted(cat.variantes.items()):
            if t == tipo:
                cats.append([tipo, var.most_common(1)[0][0], " | ".join(f"{x} ({n})" for x, n in var.most_common())])
    for f, n in fuentes.most_common():
        cats.append(["fuente de financiación", f, f"{n} registros"])
    hoja("Catálogos", ["Catálogo", "Valor que queda", "Variantes (registros)"], cats, [24, 50, 90])

    del wb["Sheet"]
    wb.save(ruta)


def guardar_usuarios():
    """Conserva correo, rol y equipo de los usuarios registrados para no perderlos al reimportar."""
    from sqlalchemy import inspect, text
    if not inspect(engine).has_table("personas"):
        return []
    with engine.connect() as c:
        return c.execute(text(
            "SELECT p.nombre, p.clave, p.email, p.rol, p.activo, e.nombre AS equipo FROM personas p "
            "LEFT JOIN equipos e ON e.id = p.equipo_id WHERE p.email IS NOT NULL")).mappings().all()


def restaurar_usuarios(s, cat, usuarios):
    for u in usuarios:
        p = cat.personas.get(u["clave"])
        if p is None:  # usuario que no aparece en el Excel (p. ej. administradores)
            p = m.Persona(nombre=u["nombre"], clave=u["clave"])
            s.add(p)
        eq = cat.equipos.get(L.clave(u["equipo"])) if u["equipo"] else None
        p.email, p.rol, p.activo, p.equipo_id = u["email"], u["rol"], u["activo"], eq.id if eq else None
    if usuarios:
        print(f"Usuarios conservados: {len(usuarios)}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--reporte", default="reporte_importacion.xlsx")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.excel, read_only=True, data_only=True)
    for h in (HOJA_CT, HOJA_CV):
        if h not in wb.sheetnames:
            sys.exit(f"No se encontró la hoja '{h}'")

    usuarios = guardar_usuarios()
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)

    log, uso_personas = L.Bitacora(), Counter()
    with SessionLocal() as s:
        cat = Catalogos(s)
        enc_ct = list(next(wb[HOJA_CT].iter_rows(max_row=1, values_only=True)))
        enc_cv = list(next(wb[HOJA_CV].iter_rows(max_row=1, values_only=True)))
        n_ct = importar_contratos(wb[HOJA_CT], s, cat, log, uso_personas, enc_ct)
        n_cv = importar_convenios(wb[HOJA_CV], s, cat, log, uso_personas, enc_cv)
        cat.nombres_canonicos()
        restaurar_usuarios(s, cat, usuarios)
        s.commit()

        from sqlalchemy import func, select
        cuenta = lambda modelo, *w: s.scalar(select(func.count()).select_from(modelo).where(*w))
        fuentes = Counter(s.scalars(select(m.Contrato.fuente_financiacion)).all())
        dup = s.execute(select(m.Contrato.numero, m.Contrato.fuente_financiacion, func.count())
                        .group_by(m.Contrato.numero, m.Contrato.fuente_financiacion)
                        .having(func.count() > 1)).all()
        for numero, fuente, n in dup:
            log.anotar(HOJA_CT, None, numero, "No CT-AÑO + FUENTE", f"{numero} / {fuente}",
                       f"número + fuente repetido en {n} filas")
        resumen = [
            ["Contratos cargados", n_ct], ["Convenios cargados", n_cv],
            ["Equipos", len(cat.equipos)], ["Estados", len(cat.estados)], ["Personas", len(cat.personas)],
            ["Asignaciones", cuenta(m.Asignacion)],
            ["Informes (contratos)", cuenta(m.Informe, m.Informe.contrato_id.isnot(None))],
            ["Informes (convenios)", cuenta(m.Informe, m.Informe.convenio_id.isnot(None))],
            ["Desembolsos", cuenta(m.Desembolso)], ["Fuentes CDR", cuenta(m.FuenteCDR)],
            ["Aliados de convenios", cuenta(m.ConvenioAliado)], ["Modificaciones", cuenta(m.Modificacion)],
            ["Alertas", cuenta(m.Alerta)],
            ["Combinaciones número + fuente repetidas", len(dup)],
            ["Celdas con problemas", len(log.problemas)],
        ]
    escribir_reporte(a.reporte, log, cat, uso_personas, resumen, fuentes)
    for k, v in resumen:
        print(f"{k:45} {v}")
    print(f"\nReporte: {a.reporte}")


if __name__ == "__main__":
    main()
